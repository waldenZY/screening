"""Standalone quality review and production-input compatibility, with simulated models."""
import contextlib
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from screening.module_tests import run_admission, run_quality_annotation as quality
from screening.module_tests.tests.test_run_admission import CONFIG, CATEGORY, feature, record, fake_client
from screening.quality.admission import build_admission_result
from screening.quality.annotation import quality_annotation_input
from screening.candidate.contract import parse_conversation


class QualityModuleTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)

    def write(self,rows,name="input.jsonl"):
        path=self.root/name
        path.write_text("".join(json.dumps(row,ensure_ascii=False)+"\n" for row in rows),encoding="utf-8")
        return path

    def test_labeled_raw_text_preserves_actual_model_markers(self):
        original="先打开相机 | [ASR] 请拍照"
        parsed=parse_conversation(original)
        row={**record(parsed["target_query"]," | ".join(parsed["context_turns"])),
            "raw_text":original,"stage1":{"admission":build_admission_result(feature(["请拍照"]),CATEGORY)["admission"]}}
        payload=run_admission.build_admission_input(row)
        self.assertEqual(payload["conversation_turns"],["先打开相机","[ASR] 请拍照"])
        candidate={"conversation":parsed,"taxonomy":{"l1":"设备","l2":"控制"}}
        admission={"admission":row["stage1"]["admission"]}
        self.assertEqual(quality.build_quality_input(row),quality_annotation_input(candidate,admission))
        wrong=copy.deepcopy(row);wrong["raw_text"]="先打开相机 | 删除照片"
        with self.assertRaises(ValueError):run_admission.build_admission_input(wrong)

    def test_admission_results_to_quality_review_preserves_rows_and_labels(self):
        source=self.write([record("帮我设闹钟"),record("=2+2"),record("打开客厅灯")])
        with contextlib.redirect_stdout(io.StringIO()):
            code=run_admission.run_admission_test(source,self.root/"admission",llm_config=CONFIG,llm_call=fake_client([]))
        self.assertEqual(code,0)
        calls=[]
        def client(messages,**kwargs):
            prompt=messages[0]["content"];calls.append(prompt)
            payload=json.loads(prompt[prompt.rfind("\n{")+1:])
            self.assertEqual(set(payload),{"sample_type","conversation_turns","candidate_category","admission","quality_rubric_version"})
            score=2 if payload["conversation_turns"][-1]=="打开客厅灯" else 5
            return json.dumps({"quality_validity":"invalid" if score==2 else "valid","quality_intent_clarity":"clear",
                "quality_overall_score":score,"quality_issue_types":["context_dependent"] if score==2 else ["none"]}),1,2
        with contextlib.redirect_stdout(io.StringIO()):
            code=quality.run_quality_test(self.root/"admission/results.jsonl",self.root/"quality",llm_config=CONFIG,llm_call=client)
        self.assertEqual(code,0);self.assertEqual(len(calls),3)
        rows=[json.loads(line) for line in (self.root/"quality/results.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual([row["source_line"] for row in rows],[1,2,3])
        self.assertEqual(rows[0]["annotation_input"]["admission"]["screening"]["route"],"clarification")
        self.assertEqual(rows[0]["quality_result"]["quality_overall_score"],5)
        self.assertEqual(rows[2]["quality_result"]["quality_overall_score"],2)
        info=json.loads((self.root/"quality/run.json").read_text(encoding="utf-8"))
        self.assertEqual(info["counts"]["failed"],0);self.assertEqual(info["counts"]["low_score"],1)
        self.assertNotIn(CONFIG.api_key,(self.root/"quality/run.json").read_text(encoding="utf-8"))
        book=load_workbook(self.root/"quality/results.xlsx")
        try:
            sheet=book.active;headers=[cell.value for cell in sheet[1]]
            self.assertEqual(sheet.max_row,4)
            cell=sheet.cell(3,headers.index("target_query")+1)
            self.assertEqual(cell.value,"=2+2");self.assertEqual(cell.data_type,"s")
            self.assertIn("quality_issue_types",headers)
        finally:book.close()

    def test_no_fabricated_admission_missing_category_and_tampered_input(self):
        source=self.write([record(),record(category=False)])
        with contextlib.redirect_stdout(io.StringIO()):
            run_admission.run_admission_test(source,self.root/"admission",llm_config=CONFIG,llm_call=fake_client([]))
        rows=[json.loads(line) for line in (self.root/"admission/results.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(quality.build_quality_input(rows[0])["candidate_category"],{"L1":"设备","L2":"控制"})
        with self.assertRaises(ValueError):quality.build_quality_input(rows[1])
        rows[0]["source_record"]["target_query"]="关掉客厅灯"
        with self.assertRaises(ValueError):quality.build_quality_input(rows[0])
        path=self.write([record(),rows[1],rows[0]],"invalid.jsonl")
        calls=[]
        with contextlib.redirect_stdout(io.StringIO()):
            code=quality.run_quality_test(path,self.root/"quality",llm_config=CONFIG,llm_call=fake_client(calls))
        self.assertEqual(code,2);self.assertEqual(calls,[])

    def test_response_failure_is_not_a_low_score_and_does_not_retry(self):
        admission=build_admission_result(feature(["打开客厅灯"]),CATEGORY)["admission"]
        row={**record(),"stage1":{"admission":admission}}
        path=self.write([row,row])
        calls=[]
        def client(messages,**kwargs):
            calls.append(messages)
            if len(calls)==1:return '{"quality_overall_score":5}',1,2
            raise RuntimeError("must not expose endpoint or secret")
        with contextlib.redirect_stdout(io.StringIO()):
            code=quality.run_quality_test(path,self.root/"quality",llm_config=CONFIG,llm_call=client)
        self.assertEqual(code,2);self.assertEqual(len(calls),2)
        rows=[json.loads(line) for line in (self.root/"quality/results.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual([r["status"] for r in rows],["response_error","request_error"])
        self.assertTrue(all(r["quality_result"] is None for r in rows))
        self.assertNotIn("must not expose",(self.root/"quality/results.jsonl").read_text(encoding="utf-8"))
        with self.assertRaises(ValueError):quality.run_quality_test(path,self.root/"quality",llm_config=CONFIG,llm_call=client)

    def test_frozen_internal_and_new_delivery_inputs_match_production_payload(self):
        from screening.module_tests.tests.test_downstream import StageBoundaryTests
        from screening.pipeline.cli import main as pipeline_main
        profile=StageBoundaryTests()._profile(self.root)
        with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
            code=pipeline_main(["--config",str(profile),"--no-publish"],llm_config_override=CONFIG,llm_call_override=fake_client([]))
        self.assertEqual(code,0)
        run=next((self.root/"outputs/runs").iterdir())
        for branch in ("P0","P1","P2"):
            freeze=[json.loads(line) for line in (run/branch/"stage1.frozen.jsonl").read_text(encoding="utf-8").splitlines()]
            internal=[json.loads(line) for line in (run/branch/"screened.canonical.jsonl").read_text(encoding="utf-8").splitlines()]
            delivered=[json.loads(line) for line in next((run/branch).glob("*_v4_labeled.jsonl")).read_text(encoding="utf-8").splitlines()]
            self.assertTrue(freeze)
            for frozen,canonical,labeled in zip(freeze,internal,delivered):
                expected=quality_annotation_input(frozen["candidate"],{"admission":frozen["quality"]["admission"]})
                self.assertEqual(quality.build_quality_input(frozen),expected)
                self.assertEqual(quality.build_quality_input(canonical),expected)
                self.assertEqual(quality.build_quality_input(labeled),expected)
                self.assertEqual(run_admission.build_admission_input(frozen),{
                    key:expected[key] for key in ("sample_type","conversation_turns","candidate_category")})


if __name__=="__main__":unittest.main()
