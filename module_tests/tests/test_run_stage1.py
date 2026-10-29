"""Admission semantics, isolated requests, replay and production round trips."""
from __future__ import annotations
import contextlib
import copy
import io
import json
import hashlib
import tempfile
import unittest
from unittest.mock import patch
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook

from screening.quality.stage1 import (DEFAULT_POLICY, QualityResponseError, build_stage1_result,
    validate_features, validate_stage1_result, screen, stage1_passes, admission_group, stage1_issue_summary, mapped_intent_label)
from screening.quality.runner import QUALITY_CONTRACT_VERSION, QualityRunner, QualityRunConfig, QualityStageIncomplete
from screening.module_tests import run_stage1 as standalone
from screening.delivery.eligibility import delivery_labels, executable_reasons
from screening.delivery.publish import collect_and_validate_artifacts, validate_manifest_artifacts, PublicationError
from screening.delivery.screened import validate_screened_canonical_record, ScreenedCanonicalError
from screening.delivery.labeled import map_screened_to_labeled, validate_labeled_record, LabeledExportError
from screening.delivery.report import validate_run_report, RunReportError
from screening.pipeline.cli import main as pipeline_main
from screening.pipeline.cli import build_parser
from screening.pipeline.config import load_pipeline_config, apply_pipeline_overrides
from screening.shared.llm import LlmConfig
from screening.shared.risk_keywords import load_risk_keywords

PACKAGE = Path(__file__).resolve().parents[2]
CONFIG = LlmConfig(api_url="https://offline.example.invalid/v1/chat/completions", api_key="offline-secret",
    model="offline-model", max_tokens=4096, timeout=3, stream=False, enable_thinking=False, skip_special_tokens=False)
CATEGORY = {"status": "reviewed", "fit": "exact"}


def feature(turns, *, intent="task", points=(), deps=()):
    return {"intent_label": intent, "intent_reason": "合成测试观察：明确的操作、回复需求或无法确定需求。",
        "open_points": [{"kind": kind, "description": "需要确定的信息"} for kind in points],
        "dependencies": [{"source": source, "description": "必要依赖及来源依据"} for source in deps]}


def record(target="打开客厅灯", context="", category=True):
    turns = (context.split(" | ") if context else []) + [target]
    return {"sample_id": "reference", "context_text": context, "target_query": target,
        "sample_type": "multi_turn" if context else "single_turn", "conversation_turns": turns,
        "candidate_category": {"L1": "设备", "L2": "控制"} if category else {"L1": None, "L2": None}}


def fake_features(turns):
    # Explicit mock responses exercise plumbing; they are not a semantic classifier or gold labels.
    target = turns[-1]
    if target in {"闹钟", "冰糖冬枣", "视频剪辑", "临沂"} or (len(turns) == 1 and target in {"嗯", "这样", "他说也不错", "下个路口左转"}):
        return feature(turns, intent="undetermined")
    if target == "帮我设闹钟":
        return feature(turns, points=("parameter",))
    if target == "总结这份文件":
        return feature(turns, intent="dialogue", deps=("missing_material",))
    if target == "选择第二个":
        return feature(turns, deps=("missing_context",))
    if target == "总结当前视频":
        return feature(turns, intent="dialogue", deps=("current_screen_media",))
    if target == "推荐几个周末去处":
        return feature(turns, intent="dialogue", points=("object",))
    if target in {"写一封请假邮件", "查北京明天天气", "怎么压缩图片", "我很难过", "总结已提供的文字", "制定周末旅行计划", "先查询附近餐厅，再比较价格并选择最近一家"}:
        return feature(turns, intent="dialogue")
    return feature(turns)


def fake_client(calls):
    def call(messages, **kwargs):
        prompt = messages[0]["content"]
        calls.append(prompt)
        if prompt.startswith("# Stage1：评测"):
            payload = json.loads(prompt[prompt.rfind("\n{")+1:])
            assert set(payload) == {"sample_type", "conversation_turns"}
            value = fake_features(payload["conversation_turns"])
        elif prompt.startswith("# Stage1：独立"):
            payload = json.loads(prompt[prompt.rfind("\n{")+1:])
            assert set(payload) == {"sample_type", "conversation_turns", "candidate_category"}
            value = {"category_fit": "exact"}
        elif prompt.startswith("# P2"):
            payload = json.loads(prompt[prompt.rfind("\n{")+1:])
            assert "stage1_features" in payload and "screening" not in payload
            assert set(payload["stage1_features"]) == {"intent_label", "intent_reason", "open_points", "dependencies"}
            value = {"is_challenge": True, "primary_challenge": "multi_constraint", "challenge_types": ["multi_constraint"]}
        elif prompt.startswith("# 第二阶段"):
            payload = json.loads(prompt[prompt.rfind("\n{")+1:])
            assert set(payload) == {"sample_type", "conversation_turns"}
            # Deliberate cross-stage disagreement must not veto direct eligibility.
            value = {"task_difficulty": "simple", "execution_issues": ["attachment_dependency"]}
        else:
            return "OK", 1.0, 2.0
        return json.dumps(value, ensure_ascii=False), 1.0, 2.0
    return call


class Stage1ContractTests(unittest.TestCase):
    def test_four_field_contract_and_removed_fields_rejected(self):
        observation=feature(["设闹钟"],points=("parameter",),deps=("system_app_state",))
        self.assertEqual(set(validate_features(observation)),{"intent_label","intent_reason","open_points","dependencies"})
        for field in ("text_state","anchor","goal","objects","constraints","evidence"):
            old=copy.deepcopy(observation);old[field]="old"
            with self.subTest(field=field), self.assertRaises(QualityResponseError): validate_features(old)
        for field in ("locator","evidence"):
            old=copy.deepcopy(observation);old["dependencies"][0][field]=None
            with self.assertRaises(QualityResponseError): validate_features(old)
        for change in ({"intent_reason":""},{"intent_reason":None},{"intent_label":"qa"}):
            old={**observation,**change}
            with self.assertRaises(QualityResponseError): validate_features(old)
        for source in ("user_to_specify","made_up",["provided","missing_context"]):
            old=copy.deepcopy(observation);old["dependencies"][0]["source"]=source
            with self.assertRaises(QualityResponseError): validate_features(old)
        old=copy.deepcopy(DEFAULT_POLICY);old["version"]=1
        with self.assertRaises(QualityResponseError): build_stage1_result(observation,CATEGORY,old)
        with self.assertRaises(QualityResponseError): mapped_intent_label(feature(["闹钟"],intent="undetermined"))

    def test_five_groups_and_dialogue_open_points_direct(self):
        features=[feature(["a"],intent="dialogue"),feature(["b"],intent="dialogue",points=("object",)),
                  feature(["c"]),feature(["d"],points=("parameter",)),feature(["e"],intent="undetermined")]
        expected=["dialogue_without_open_points","dialogue_with_open_points","task_direct","task_clarification","excluded"]
        values=[build_stage1_result(f,CATEGORY) for f in features]
        self.assertEqual([admission_group(v) for v in values],expected)
        self.assertEqual([v["stage1"]["screening"]["route"] for v in values],["direct","direct","direct","clarification","excluded"])
        summary=stage1_issue_summary(values)
        self.assertEqual(summary["stage1_admission_groups"],dict.fromkeys(expected,1))
        self.assertEqual(summary["stage1_intent_counts"],{"dialogue":2,"task":2,"undetermined":1})

    def test_versioned_illustrative_fixtures(self):
        fixture=json.loads((PACKAGE/"artifacts/evals/stage1_cases.json").read_text(encoding="utf-8"))
        self.assertEqual(fixture["quality_contract_version"],QUALITY_CONTRACT_VERSION)
        for case in fixture["cases"]:
            with self.subTest(case=case["id"]):
                result=build_stage1_result(case["features"],case["category"],case["policy"])
                self.assertEqual(result["stage1"]["screening"],case["expected_screening"])

    def test_ordered_rules_and_independent_category(self):
        for target, expected in (("打开客厅灯","direct"),("闹钟","excluded"),("视频剪辑","excluded"),
                ("冰糖冬枣","excluded"),("临沂","excluded"),("帮我设闹钟","clarification"),
                ("嗯","excluded"),("这样","excluded"),("他说也不错","excluded"),("下个路口左转","excluded"),
                ("总结这份文件","excluded"),("选择第二个","excluded"),("总结当前视频","direct"),
                ("推荐几个周末去处","direct")):
            with self.subTest(target=target):
                result=build_stage1_result(fake_features([target]),CATEGORY)
                self.assertEqual(result["stage1"]["screening"]["route"],expected)
                self.assertEqual(stage1_passes(result),expected!="excluded")
        f=feature(["总结文件并设闹钟"],points=("parameter",),deps=("missing_material","missing_context"))
        result=build_stage1_result(f,{"status":"reviewed","fit":"partial"})
        self.assertEqual(result["stage1"]["screening"]["reasons"],["missing_material","missing_context","category_partial"])
        for intent in ("dialogue","task"):
            for source in ("missing_material","missing_context","unknown"):
                result=build_stage1_result(feature(["a"],intent=intent,points=("parameter",),deps=(source,)),CATEGORY)
                self.assertEqual(result["stage1"]["screening"]["route"],"excluded")

    def test_replay_policy_and_invalid_observation_or_decision(self):
        f=fake_features(["帮我设闹钟"]);default=build_stage1_result(f,CATEGORY)
        policy=copy.deepcopy(DEFAULT_POLICY);policy["allowed_routes"]=["direct"]
        self.assertFalse(stage1_passes(build_stage1_result(f,CATEGORY,policy)))
        qa=feature(["推荐"],intent="dialogue",points=("object",))
        self.assertTrue(stage1_passes(build_stage1_result(qa,CATEGORY,policy)))
        bad=copy.deepcopy(default);bad["stage1"]["screening"]["admitted"]=False
        with self.assertRaises(QualityResponseError): validate_stage1_result(bad)
        runtime=feature(["查天气并读取日历"],intent="dialogue",deps=("public_information","system_app_state"))
        self.assertTrue(stage1_passes(build_stage1_result(runtime,CATEGORY)))
        policy["runtime_capabilities"]=[]
        self.assertIn("runtime_unavailable",screen(runtime,CATEGORY,policy)["reasons"])
        f["intent_label"]="undetermined"
        self.assertEqual(screen(f,CATEGORY,DEFAULT_POLICY)["route"],"excluded")

    def test_missing_category_does_not_change_text_route(self):
        result=build_stage1_result(fake_features(["推荐几个周末去处"]),{"status":"not_provided","fit":None})
        self.assertEqual(result["stage1"]["screening"]["route"],"direct")
        self.assertEqual(result["stage1"]["screening"]["reasons"],["category_not_provided"])
        self.assertEqual(admission_group(result),"excluded")

    def test_executable_and_intent_from_stage1(self):
        with tempfile.TemporaryDirectory() as td:
            risk_path=Path(td)/"risk.xlsx"
            workbook=Workbook();workbook.active.title="Sheet1";workbook.active.append(["关键词","风险级别"]);workbook.active.append(["关机","高"]);workbook.save(risk_path)
            risk=load_risk_keywords(risk_path)
            for turns,expected in ((["打开客厅灯"],1),(["暂停"],0),(["帮我设闹钟"],0),
                    (["请关机"],0),(["开灯","调暗一点"],0),(["推荐几个周末去处"],1),(["写一封请假邮件"],1)):
                candidate={"conversation":{"target_query":turns[-1],"sample_type":"multi_turn" if len(turns)>1 else "single_turn"}}
                stage1=build_stage1_result(fake_features(turns),CATEGORY)
                labels=delivery_labels(candidate,stage1,risk)
                self.assertEqual(labels["executable"],expected)
                self.assertEqual(labels["intent_label"],"qa" if stage1["stage1"]["features"]["intent_label"]=="dialogue" else "task")
                self.assertEqual(bool(labels["executable_reasons"]),not expected)
            rejected=build_stage1_result(fake_features(["闹钟"]),CATEGORY)
            self.assertIn("not_admitted",executable_reasons({"conversation":{"target_query":"闹钟","sample_type":"single_turn"}},rejected,risk))


class Stage1ModuleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup); self.root = Path(self.temp.name)

    def write(self, rows, name="input.jsonl"):
        path = self.root/name; path.write_text("".join(json.dumps(r,ensure_ascii=False)+"\n" for r in rows),encoding="utf-8"); return path

    def run_samples(self, rows, **kwargs):
        path = self.write(rows); calls=[]
        with contextlib.redirect_stdout(io.StringIO()):
            code = standalone.run_stage1_test(path,self.root/"out",llm_config=CONFIG,llm_call=fake_client(calls),**kwargs)
        results=[json.loads(x) for x in (self.root/"out/results.jsonl").read_text(encoding="utf-8").splitlines()]
        return code,calls,results

    def test_reference_duplicates_category_skip_and_replay_without_credentials(self):
        code,calls,rows=self.run_samples([record("帮我设闹钟",category=False),record("帮我设闹钟",category=False),record()])
        self.assertEqual(code,0);self.assertEqual(len(calls),4)
        self.assertEqual([r["route"] for r in rows],["clarification","clarification","direct"])
        self.assertEqual(rows[0]["stage_status"]["category"],"not_provided")
        self.assertEqual([r["source_line"] for r in rows],[1,2,3])
        policy=copy.deepcopy(DEFAULT_POLICY);policy["allowed_routes"]=["direct"]
        with contextlib.redirect_stdout(io.StringIO()):
            code=standalone.run_stage1_test(self.root/"out/results.jsonl",self.root/"replay",replay=True,policy=policy)
        self.assertEqual(code,0)
        again=[json.loads(x) for x in (self.root/"replay/results.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(again[0]["features"],rows[0]["features"])
        self.assertIn("route_disabled",again[0]["stage1_result"]["stage1"]["screening"]["reasons"])
        self.assertEqual(again[0]["observation_provenance"],rows[0]["observation_provenance"])
        workbook=load_workbook(self.root/"out/results.xlsx");self.assertEqual(workbook.active.max_row,4)
        tampered=copy.deepcopy(rows[2]);tampered["source_record"]["candidate_category"]["L2"]="其他"
        tampered["stage1_input"]["candidate_category"]["L2"]="其他"
        tampered_input=self.write([tampered],"tampered.jsonl")
        with contextlib.redirect_stdout(io.StringIO()):
            code=standalone.run_stage1_test(tampered_input,self.root/"bad-replay",replay=True)
        self.assertEqual(code,2)


    def test_invalid_inputs_and_output_directory(self):
        bad=record();bad["conversation_turns"]=["不同输入"]
        code,calls,rows=self.run_samples([bad,record("打开客厅灯", "旧话题")])
        self.assertEqual(code,2);self.assertEqual(rows[0]["status"],"input_error");self.assertIsNone(rows[0]["stage1_pass"])
        self.assertEqual(len(calls),2)
        with self.assertRaises(ValueError): standalone.run_stage1_test(self.root/"input.jsonl",self.root/"out",llm_config=CONFIG)
        old=record(); old.update(text="打开客厅灯",turn_count=1,L1="设备",L2="控制",intent_label="old-label")
        payload=standalone.build_stage1_input(old);self.assertNotIn("intent_label",payload)

    def test_category_failure_preserves_successful_features(self):
        calls=[];good=fake_client(calls)
        def client(messages,**kwargs):
            if messages[0]["content"].startswith("# Stage1：独立"): return "not json",0,0
            return good(messages,**kwargs)
        with contextlib.redirect_stdout(io.StringIO()):
            code=standalone.run_stage1_test(self.write([record()]),self.root/"out",llm_config=CONFIG,llm_call=client)
        row=json.loads((self.root/"out/results.jsonl").read_text(encoding="utf-8"))
        self.assertEqual(code,2);self.assertEqual(row["status"],"response_error");self.assertIsNotNone(row["features"]);self.assertIsNone(row["stage1_pass"])
        counts=json.loads((self.root/"out/run.json").read_text(encoding="utf-8"))["counts"]
        self.assertEqual(counts["failed"],1)
        self.assertEqual(sum(counts["stage1_admission_groups"].values()),0)

    def test_engineering_failure_not_exclusion_and_readable_fields(self):
        calls=[];good=fake_client(calls)
        def client(messages,**kwargs):
            if messages[0]["content"].startswith("# Stage1：评测"):
                payload=json.loads(messages[0]["content"].split("## 待识别输入\n")[1])
                if payload["conversation_turns"]==["非法模型响应"]: return "not json",0,0
            return good(messages,**kwargs)
        with contextlib.redirect_stdout(io.StringIO()):
            code=standalone.run_stage1_test(self.write([record("推荐几个周末去处"),record("闹钟"),record("非法模型响应")]),self.root/"out",llm_config=CONFIG,llm_call=client)
        self.assertEqual(code,2)
        counts=json.loads((self.root/"out/run.json").read_text(encoding="utf-8"))["counts"]
        self.assertEqual((counts["passed"],counts["rejected"],counts["failed"]),(1,1,1))
        self.assertEqual(sum(counts["stage1_admission_groups"].values()),2)
        self.assertEqual(counts["stage1_admission_groups"]["dialogue_with_open_points"],1)
        workbook=load_workbook(self.root/"out/results.xlsx")
        headers=[c.value for c in workbook.active[1]]
        for name in ("intent_label","intent_reason","open_points","dependencies","admission_group"):
            self.assertIn(name,headers)
        self.assertEqual(workbook.active.cell(2,headers.index("intent_label")+1).value,"dialogue")
        self.assertEqual(workbook.active.max_row,4)
        workbook.close()

    def test_replay_old_observation_rejected_without_credentials(self):
        code,calls,rows=self.run_samples([record()]);self.assertEqual(code,0)
        old=rows[0]
        old_features={"text_state":"recoverable","anchor":"task","goal":"打开客厅灯","objects":[],"constraints":[],"open_points":[],"dependencies":[]}
        old["features"]=old_features;old["stage1_result"]["stage1"]["features"]=old_features
        old_input=self.write([old],"old.jsonl")
        with contextlib.redirect_stdout(io.StringIO()):
            code=standalone.run_stage1_test(old_input,self.root/"old-replay",replay=True)
        self.assertEqual(code,2)
        row=json.loads((self.root/"old-replay/results.jsonl").read_text(encoding="utf-8"))
        self.assertIsNone(row["stage1_pass"])
        self.assertEqual(row["status"],"response_error")

    def test_local_reference_shape_when_available(self):
        path=PACKAGE.parent/"当前透传到外部的样本仅供参考/query_samples_all.jsonl"
        if not path.exists(): self.skipTest("本地参考文件不随代码入库")
        rows=[json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
        self.assertEqual(len(rows),202)
        payloads=[standalone.build_stage1_input(r) for r in rows]
        self.assertEqual(sum(p["sample_type"]=="single_turn" for p in payloads),149)
        self.assertEqual(sum(p["sample_type"]=="multi_turn" for p in payloads),53)
        self.assertEqual(sum(n>1 for n in Counter(tuple(p["conversation_turns"]) for p in payloads).values()),12)
        self.assertTrue(all(p["candidate_category"]=={"L1":None,"L2":None} for p in payloads))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),"4cd04d2bf81d5aa7b25d9b31e254fc309944ab2ea44c718f9299c41457b11510")


class OptionalConfigurationTests(unittest.TestCase):
    def test_default_disabled_overrides_and_invalid_type(self):
        config=load_pipeline_config(PACKAGE/"artifacts/config/pipeline.default.json")
        self.assertFalse(config.stage2_enabled)
        self.assertIsNone(config.enabled_stage2_prompt)
        enabled=apply_pipeline_overrides(config,{"stage2_enabled":True})
        self.assertTrue(enabled.stage2_enabled)
        self.assertEqual(enabled.enabled_stage2_prompt,PACKAGE/"quality/optional_stage2/prompt.md")
        disabled=apply_pipeline_overrides(enabled,{"stage2_enabled":False,"stage2_prompt":"does-not-exist.md"})
        self.assertIsNone(disabled.enabled_stage2_prompt)
        self.assertTrue(build_parser().parse_args(["--enable-stage2"]).stage2_enabled)
        self.assertFalse(build_parser().parse_args(["--disable-stage2"]).stage2_enabled)
        with tempfile.TemporaryDirectory() as td:
            profile=json.loads((PACKAGE/"artifacts/config/pipeline.default.json").read_text(encoding="utf-8"))
            profile["optional_stage2"]["enabled"]="false"
            path=Path(td)/"bad-config.json";path.write_text(json.dumps(profile),encoding="utf-8")
            with self.assertRaises(ValueError): load_pipeline_config(path)


class RunnerIsolationTests(unittest.TestCase):
    def test_component_cache_category_isolation_retry_and_restart(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            calls=[]
            client=fake_client(calls)
            attempts={"category":0}
            def flaky(messages,**kwargs):
                if messages[0]["content"].startswith("# Stage1：独立"):
                    attempts["category"]+=1
                    if attempts["category"]==1:
                        raise RuntimeError("do not log endpoint or credential here")
                return client(messages,**kwargs)
            kwargs=dict(stage1_prompt=standalone.DEFAULT_PROMPT_PATH.read_text(encoding="utf-8"),
                category_prompt=standalone.DEFAULT_CATEGORY_PROMPT_PATH.read_text(encoding="utf-8"),
                challenge_prompt=(PACKAGE/"artifacts/prompts/quality_challenge.md").read_text(encoding="utf-8"),
                model=CONFIG.model,llm_options=CONFIG.options(),config=QualityRunConfig(workers=2,sleep_seconds=0,max_retries=1))
            from screening.candidate.contract import parse_conversation
            conversation=parse_conversation("打开客厅灯")
            a={"source_identity_sha256":"a"*64,"conversation_sha256":"c"*64,
                "conversation":conversation,"source":{"source_line":1},"taxonomy":{"l1":"设备","l2":"控制","l2_key":"设备/控制"}}
            b=copy.deepcopy(a);b["source_identity_sha256"]="b"*64;b["source"]["source_line"]=2;b["taxonomy"]["l2"]="其他"
            runner=QualityRunner(root,llm_call=flaky,**kwargs)
            with self.assertRaises(QualityStageIncomplete): runner.run_stage1([a])
            self.assertNotIn("do not log endpoint",(root/"stage1_category_attempts.jsonl").read_text(encoding="utf-8"))
            self.assertEqual(len([p for p in calls if p.startswith("# Stage1：评测")]),1)
            # A fresh runner resumes the feature cache after category failure.
            restarted=QualityRunner(root,llm_call=client,**kwargs)
            first=restarted.run_stage1([a,b])
            self.assertEqual(set(first),{"a"*64,"b"*64})
            self.assertEqual(len([p for p in calls if p.startswith("# Stage1：评测")]),1)
            self.assertEqual(len([p for p in calls if p.startswith("# Stage1：独立")]),2)
            before=len(calls);restarted.run_stage1([a,b]);self.assertEqual(len(calls),before)
            self.assertNotEqual(restarted.task_id("stage1_category",a),restarted.task_id("stage1_category",b))
            self.assertEqual(restarted.task_id("stage1_features",a),restarted.task_id("stage1_features",b))
            policy=copy.deepcopy(DEFAULT_POLICY);policy["allowed_routes"]=["direct"]
            with self.assertRaises(ValueError): QualityRunner(root,llm_call=client,screening_policy=policy,**kwargs)
            contract_path=root/"quality_run_contract.json"
            old=json.loads(contract_path.read_text(encoding="utf-8"));old["quality_contract_version"]=11
            contract_path.write_text(json.dumps(old),encoding="utf-8")
            with self.assertRaises(ValueError): QualityRunner(root,llm_call=client,**kwargs)


class ProductionRoundTripTests(unittest.TestCase):
    def test_pipeline_resume_export_publish_and_policy_tamper(self):
        # A missing optional module must not affect the default chain.
        with patch.dict("sys.modules", {"screening.quality.optional_stage2": None}):
            self._roundtrip(False)

    @unittest.skipUnless((PACKAGE/"quality/optional_stage2").is_dir(), "可选模块已移除")
    def test_optional_stage2_preserves_identical_delivery_fields(self):
        self._roundtrip(True)

    def _roundtrip(self, stage2_enabled):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/"cluster.jsonl";taxonomy=root/"taxonomy.json";risk=root/"risk.xlsx"
            texts=["打开客厅灯","帮我设闹钟","冰糖冬枣","总结当前视频","请关机","暂停",
                "把音量降低一点","给妈妈拨打电话","播放音乐 | 换成古典音乐", "先查询附近餐厅，再比较价格并选择最近一家",
                "明早六点叫我 | 改成七点", "先查询天气，然后把结果发送给妈妈", "添加日历 | 添加桌面日历 | 嗯，是的",
                "总结这份文件", "选择第二个", "推荐几个周末去处", "写一封请假邮件"]
            source.write_text("".join(json.dumps({"session":f"s{i}","text":text,"cluster":1,"cluster_prob":0.4 if i%2 else 0.95,"frequency":i+1},ensure_ascii=False)+"\n" for i,text in enumerate(texts)),encoding="utf-8")
            taxonomy.write_text(json.dumps({"categories":[{"name":"设备","total_size":len(texts),"assigned_clusters":[1],"sub_intents":[{"name":"控制","size":len(texts),"cluster_ids":[1]}]}]},ensure_ascii=False),encoding="utf-8")
            workbook=Workbook();workbook.active.title="Sheet1";workbook.active.append(["关键词","风险级别"]);workbook.active.append(["关机","高"]);workbook.save(risk)
            profile=json.loads((PACKAGE/"artifacts/config/pipeline.default.json").read_text(encoding="utf-8"))
            profile["inputs"]={"source":str(source),"taxonomy":str(taxonomy),"risk_keywords":str(risk)}
            profile["optional_stage2"]={"enabled":stage2_enabled,"prompt":str(PACKAGE/"quality/optional_stage2/prompt.md") if stage2_enabled else str(root/"nonexistent-stage2.md")}
            profile["outputs"]["output_root"]=str(root/"outputs")
            for k in profile["prompts"]: profile["prompts"][k]=str((PACKAGE/"artifacts/config"/profile["prompts"][k]).resolve())
            profile["quality"]={"workers":3,"max_retries":2,"sleep_seconds":0}
            profile["index"].update(min_free_space_gb=0,reserve_free_space_gb=0,sqlite_cache_mb=1)
            profile["sampling"]["p1"]["target_n"]=14;profile["sampling"]["p2"].update(target_n=3,batch_size=6,soft_min_by_primary={k:0 for k in profile["sampling"]["p2"]["soft_min_by_primary"]})
            config_path=root/"pipeline.json";config_path.write_text(json.dumps(profile,ensure_ascii=False),encoding="utf-8")
            calls=[];client=fake_client(calls)
            errors = io.StringIO()
            output = io.StringIO()
            with contextlib.redirect_stdout(output),contextlib.redirect_stderr(errors):
                code=pipeline_main(["--config",str(config_path),"--publish-root",str(root/"live-delivery")],llm_config_override=CONFIG,llm_call_override=client)
            self.assertEqual(code,0, output.getvalue() + errors.getvalue()[-1200:])
            run=next((root/"outputs/runs").iterdir());manifest=json.loads((run/"manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"],"published")
            self.assertEqual(manifest["contract"]["labeled_mapping_version"],9)
            self.assertNotIn("legacy_mapping_version",manifest["contract"])
            self.assertEqual(len(list((root/"live-delivery").rglob("*_v8_labeled.jsonl"))),3)
            validate_manifest_artifacts(run,manifest)
            originals={r["session"]:r for r in (json.loads(line) for line in source.read_text(encoding="utf-8").splitlines())}
            all_records=[]
            for name in ("P0","P1","P2"):
                rows=[json.loads(x) for x in (run/name/"screened.canonical.jsonl").read_text(encoding="utf-8").splitlines()]
                self.assertTrue(rows);all_records+=rows
                exported=next((run/name).glob("*_v8_labeled.jsonl"))
                delivered=[json.loads(x) for x in exported.read_text(encoding="utf-8").splitlines()]
                self.assertTrue(all(len(r)==14 and "stage1" in r and r["intent_label"]=={"dialogue":"qa","task":"task"}[r["stage1"]["features"]["intent_label"]] for r in delivered))
                for item,internal in zip(delivered,rows):
                    original=originals[item["session_id"]]
                    self.assertEqual(item["frequency"],original["frequency"])
                    self.assertEqual(item["cluster_prob"],original["cluster_prob"])
                    self.assertEqual(item["cluster_id"],original["cluster"])
                    self.assertEqual(item["stage1"],internal["quality"]["stage1"])
            self.assertTrue(any(r["quality"]["stage1"]["screening"]["route"]=="clarification" for r in all_records))
            self.assertTrue(any(r["quality"]["stage1"]["features"]["intent_label"]=="dialogue" and r["quality"]["stage1"]["features"]["open_points"] and r["delivery_labels"]["executable"]==1 for r in all_records))
            wrong=copy.deepcopy(all_records[0]);wrong_intent="qa" if wrong["quality"]["intent_label"]=="task" else "task"
            wrong["quality"]["intent_label"]=wrong_intent;wrong["delivery_labels"]["intent_label"]=wrong_intent
            with self.assertRaises(ScreenedCanonicalError): validate_screened_canonical_record(wrong)
            wrong_export=map_screened_to_labeled(all_records[0]);wrong_export["intent_label"]=wrong_intent
            with self.assertRaises(LabeledExportError): validate_labeled_record(all_records[0]["set"],wrong_export)
            wrong_report=json.loads((run/"screening_run_report.json").read_text(encoding="utf-8"))
            wrong_report["sets"]["P0"]["stage1_admission_groups"]["task_direct"]+=1
            with self.assertRaises(RunReportError): validate_run_report(wrong_report)
            self.assertTrue(all(set(r["quality"])=={"stage1","intent_label"} and "stage2_task_id" not in r["selection"] for r in all_records))
            schemas={tuple(row) for name in ("P0","P1","P2") for row in [map_screened_to_labeled(r) for r in all_records if r["set"]==name]}
            self.assertEqual(len(schemas),1)
            self.assertNotIn("quality_task_difficulty",next(iter(schemas)))
            self.assertEqual(any(p.startswith("# 第二阶段") for p in calls),stage2_enabled)
            self.assertEqual((run/"quality/stage2_tasks.jsonl").exists(),stage2_enabled)
            self.assertEqual((run/"prompts/stage2.md").exists(),stage2_enabled)
            self.assertNotIn("quality_stage2",manifest["stages"])
            self.assertFalse(any((run/name/"checkpoints/stage2.json").exists() for name in ("P0","P1","P2")))
            if stage2_enabled:
                from screening.quality.optional_stage2 import validate_branch
                for name in ("P0","P1","P2"):
                    validate_branch(run,name,[r for r in all_records if r["set"]==name])
                sidecar=run/"P0/optional_annotations/stage2.jsonl";before=sidecar.read_bytes()
                bad=[json.loads(line) for line in before.decode("utf-8").splitlines()];bad[0]["result"]["task_difficulty"]="extreme"
                sidecar.write_text("".join(json.dumps(r,ensure_ascii=False)+"\n" for r in bad),encoding="utf-8")
                with self.assertRaises((PublicationError,QualityResponseError)): collect_and_validate_artifacts(run)
                sidecar.write_bytes(before)
            else:
                self.assertFalse(any((run/name/"optional_annotations").exists() for name in ("P0","P1","P2")))
            count=len(calls)
            with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
                code=pipeline_main(["--config",str(config_path),"--publish-root",str(root/"delivery")],llm_config_override=CONFIG,llm_call_override=client)
            self.assertEqual(code,0);self.assertEqual(len(calls),count+1) # connectivity preflight only
            self.assertEqual(len(list((root/"delivery").rglob("*_v8_labeled.jsonl"))),3)
            if stage2_enabled:
                with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
                    changed=pipeline_main(["--config",str(config_path),"--run-id",manifest["run_id"],"--disable-stage2"],llm_config_override=CONFIG,llm_call_override=client)
                self.assertEqual(changed,1)  # A different annotation mode cannot resume the same run.
            feature_rows=run/"quality/stage1_features_results.jsonl";original=feature_rows.read_text(encoding="utf-8")
            modified=[json.loads(x) for x in original.splitlines()];modified[0]["result"]["intent_label"]="undetermined";feature_rows.write_text("".join(json.dumps(r,ensure_ascii=False)+"\n" for r in modified),encoding="utf-8")
            with self.assertRaises((PublicationError,QualityResponseError)): collect_and_validate_artifacts(run)
            feature_rows.write_text(original,encoding="utf-8")
            old=json.loads(config_path.read_text(encoding="utf-8"));old["config_version"]=9;config_path.write_text(json.dumps(old),encoding="utf-8")
            with contextlib.redirect_stderr(io.StringIO()): self.assertEqual(pipeline_main(["--config",str(config_path)],llm_config_override=CONFIG,llm_call_override=client),1)


if __name__ == "__main__": unittest.main()
