"""Shareable spreadsheet view of standalone 准入评估 results."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

REVIEW_COLUMNS = (
    "source_line", "target_query", "L1", "L2", "admission_pass", "status",
    "route", "admission_group", "intent_label", "intent_reason", "open_points", "dependencies",
    "category_status", "category_fit", "screening_reasons", "features", "policy_sha256",
)
DETAIL_COLUMNS = (
    "error_type", "error_message", "elapsed_ms", "first_token_latency_ms", "full_latency_ms",
    "admission_input", "raw_response", "source_text",
)
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_LONG_TEXT_NOTE = "\n[超出 Excel 单元格长度限制，完整内容见 results.jsonl]"


def validate_xlsx_dependency() -> None:
    try:
        import openpyxl  # noqa: F401
    except ModuleNotFoundError as exc:
        raise RuntimeError("缺少 openpyxl 依赖；请在 moto 环境安装 requirements.txt") from exc


def write_annotation_results_xlsx(path: Path, results: list[dict[str, Any]]) -> None:
    """A separate review sheet for the four downstream fields and their evidence."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from ..quality.annotation import QUALITY_FIELDS

    headers = ["source_line", "status", "target_query", "context_text", "L1", "L2", *QUALITY_FIELDS,
        "admission_route", "admission_pass", "input_sha256", "error_type", "error_message",
        "elapsed_ms", "first_token_latency_ms", "full_latency_ms", "annotation_input", "raw_response", "source_record", "source_text"]
    book = Workbook();sheet = book.active;sheet.title = "下游质量标注"
    sheet.freeze_panes = "C2";sheet.sheet_view.showGridLines = False
    rows = [headers]
    for result in results:
        payload = result.get("annotation_input") or {};turns = payload.get("conversation_turns", [])
        category = payload.get("candidate_category", {});admission = payload.get("admission", {})
        quality = result.get("quality_result") or {};error = result.get("error") or {}
        provenance = result.get("observation_provenance") or {}
        admitted = admission.get("screening", {}).get("admitted")
        rows.append([result["source_line"],result["status"],turns[-1] if turns else None,
            " | ".join(turns[:-1]) if turns else None,category.get("L1"),category.get("L2"),
            *(quality.get(key) for key in QUALITY_FIELDS),admission.get("screening", {}).get("route"),
            int(admitted) if isinstance(admitted,bool) else None,provenance.get("input_sha256"),
            error.get("type"),error.get("message"),result["elapsed_ms"],result["first_token_latency_ms"],
            result["full_latency_ms"],payload,result["raw_response"],result["source_record"],result["source_text"]])
    for row_number, values in enumerate(rows,1):
        for column_number, value in enumerate(values,1):
            cell=sheet.cell(row_number,column_number,_cell_value(value))
            if isinstance(cell.value,str):
                cell.data_type="s";cell.number_format="@"
            else:
                cell.number_format="0.00" if headers[column_number-1].endswith("_ms") else "0"
            cell.font=Font(name="Arial",size=10,color="243B53")
            cell.alignment=Alignment(vertical="top",wrap_text=True)
        sheet.row_dimensions[row_number].height=34 if row_number==1 else 48
    for cell in sheet[1]:
        cell.fill=PatternFill("solid",fgColor="243B53");cell.font=Font(name="Arial",size=10,bold=True,color="FFFFFF")
    for index,name in enumerate(headers,1):
        sheet.column_dimensions[get_column_letter(index)].width = 52 if name in {"target_query","context_text","annotation_input","raw_response","source_record","error_message"} else 24
    sheet.auto_filter.ref=sheet.dimensions
    if results:
        sheet.conditional_formatting.add(f"B2:B{sheet.max_row}",FormulaRule(formula=['B2<>"ok"'],fill=PatternFill("solid",fgColor="FDE9E7")))
        column=get_column_letter(headers.index("quality_overall_score")+1)
        sheet.conditional_formatting.add(f"{column}2:{column}{sheet.max_row}",FormulaRule(
            formula=[f"AND(ISNUMBER({column}2),{column}2<5)"],fill=PatternFill("solid",fgColor="FFF2CC")))
    try:
        book.save(path)
    finally:
        book.close()


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if type(value) is int and abs(value) > 999_999_999_999_999:
        value = str(value)  # Excel only preserves 15 significant decimal digits.
    if isinstance(value, str):
        value = _XML_INVALID.sub(lambda match: f"\\u{ord(match.group()):04x}", value)
        if len(value) > 32767:
            value = value[:32767 - len(_LONG_TEXT_NOTE)] + _LONG_TEXT_NOTE
    return value


def write_results_xlsx(path: Path, results: list[dict[str, Any]]) -> None:
    """Keep one row per result; JSONL remains the complete, unmodified record."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    source_columns = list(dict.fromkeys(
        key
        for result in results if isinstance(result["source_record"], dict)
        for key in result["source_record"]
        if key not in {"target_query", "L1", "L2"}
    ))
    reserved = set(REVIEW_COLUMNS) | set(DETAIL_COLUMNS)
    source_headers = []
    for key in source_columns:
        header = key
        while header in reserved:
            header = f"source.{header}"
        reserved.add(header)
        source_headers.append(header)
    headers = [*REVIEW_COLUMNS, *source_headers, *DETAIL_COLUMNS]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "准入评估结果"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "E2"

    rows = [headers]
    for result in results:
        source = result["source_record"] if isinstance(result["source_record"], dict) else {}
        quality = (result["admission_result"] or {}).get("admission", {})
        screening = quality.get("screening", {})
        category = result.get("category_review") or {}
        supplied = source.get("candidate_category") or {}
        error = result["error"] or {}
        decision = result["admission_pass"]
        features = result.get("features") or {}
        rows.append([
            result["source_line"], source.get("target_query"), source.get("L1", supplied.get("L1")), source.get("L2", supplied.get("L2")),
            None if decision is None else int(decision), result["status"],
            result.get("route"), result.get("admission_group"), features.get("intent_label"),
            features.get("intent_reason"), features.get("open_points"), features.get("dependencies"),
            category.get("status"), category.get("fit"),
            screening.get("reasons"), result.get("features"), screening.get("policy_sha256"),
            *(source.get(key) for key in source_columns),
            error.get("type"), error.get("message"), result["elapsed_ms"],
            result["first_token_latency_ms"], result["full_latency_ms"],
            result["admission_input"], result["raw_response"], result["source_text"],
        ])

    for row_number, values in enumerate(rows, start=1):
        for column_number, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column_number, _cell_value(value))
            if isinstance(cell.value, str):
                # Queries/responses beginning with '=' or Excel error tokens are literal data.
                cell.data_type = "s"
                cell.number_format = "@"
            else:
                cell.number_format = "0.00" if headers[column_number - 1].endswith("_ms") else "0"
            cell.font = Font(name="Arial", size=10, color="243B53")
            cell.alignment = Alignment(
                horizontal="left" if isinstance(cell.value, str) else "right",
                vertical="top", wrap_text=True,
            )
        sheet.row_dimensions[row_number].height = 42 if row_number > 1 else 34

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "source_line": 12, "target_query": 48, "L1": 16, "L2": 18, "admission_pass": 14,
        "status": 20, "category_fit": 18, "category_status": 20, "route": 20,
        "features": 65, "screening_reasons": 35, "context_text": 48, "text": 48,
        "intent_reason": 48, "open_points": 45, "dependencies": 55, "admission_group": 32,
        "raw_response": 55, "admission_input": 55, "source_text": 48, "error_message": 48,
    }
    for index, name in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 22)
    sheet.auto_filter.ref = sheet.dimensions
    if results:
        sheet.conditional_formatting.add(
            f"F2:F{sheet.max_row}", FormulaRule(
                formula=['F2<>"ok"'], fill=PatternFill("solid", fgColor="FDE9E7"),
            ),
        )
        sheet.conditional_formatting.add(
            f"E2:E{sheet.max_row}", FormulaRule(
                formula=["AND(ISNUMBER(E2),E2=0)"], fill=PatternFill("solid", fgColor="FFF2CC"),
            ),
        )
    try:
        workbook.save(path)
    finally:
        workbook.close()
