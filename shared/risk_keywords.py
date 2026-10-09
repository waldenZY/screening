"""Frozen, deterministic risk-keyword rules for labeled v5."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook


RULE_VERSION = 1
PREFIXES = ("显示", "关闭", "开启", "关掉", "打开", "设置", "定时", "是否", "忽略", "双指关节")
SUFFIXES = ("模式", "检测", "开启", "关闭", "区域", "时长间隔", "选项关闭", "守护模式")
SPECIAL = {
    "高对比度文字开启": ("高对比度",),
    "缓存是否命中检测": ("缓存命中",),
    "显示重绘刷新区域": ("重绘",),
    "过渡动画缩放": ("动画缩放",),
    "响铃时长间隔": ("响铃时长",),
    "定时开关机": ("开关机", "关机"),
}


def _core_terms(keyword: str) -> tuple[str, ...]:
    if keyword in SPECIAL:
        return SPECIAL[keyword]
    core = keyword
    changed = True
    while changed:
        changed = False
        for prefix in PREFIXES:
            if core.startswith(prefix) and len(core) > len(prefix):
                core = core[len(prefix):]
                changed = True
                break
    changed = True
    while changed:
        changed = False
        for suffix in SUFFIXES:
            if core.endswith(suffix) and len(core) > len(suffix):
                core = core[:-len(suffix)]
                changed = True
                break
    return (core,)


def _digest(value: Any) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def load_risk_keywords(path: Path) -> dict[str, Any]:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"风险关键词文件不存在: {path}")
    source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("风险关键词文件不是可读取的 Excel") from exc
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError("风险关键词 Excel 缺少 Sheet1")
        rows = workbook["Sheet1"].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(header[:2]) != ("关键词", "风险级别") or any(value is not None for value in header[2:]):
            raise ValueError("风险关键词 Sheet1 必须只有 关键词、风险级别 两列")
        entries: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows, start=2):
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue
            if any(value is not None and str(value).strip() for value in row[2:]):
                raise ValueError(f"风险关键词 Sheet1 第 {row_number} 行包含额外列")
            keyword, level = row[:2]
            if not isinstance(keyword, str) or not keyword.strip() or not isinstance(level, str) or not level.strip():
                raise ValueError(f"风险关键词 Sheet1 第 {row_number} 行字段非法")
            keyword = keyword.strip()
            entries.append({"keyword": keyword, "level": level.strip(), "terms": list(dict.fromkeys((*_core_terms(keyword), keyword)))})
        if not entries:
            raise ValueError("风险关键词 Sheet1 不能为空")
    finally:
        workbook.close()
    rules = {"rule_version": RULE_VERSION, "prefixes": list(PREFIXES), "suffixes": list(SUFFIXES),
             "special": {key: list(value) for key, value in SPECIAL.items()}, "entries": entries,
             "source_sha256": source_sha256}
    rules["rules_sha256"] = _digest({key: value for key, value in rules.items() if key != "rules_sha256"})
    return rules


def validate_risk_snapshot(snapshot: Mapping[str, Any]) -> None:
    if snapshot.get("rule_version") != RULE_VERSION or snapshot.get("prefixes") != list(PREFIXES) or snapshot.get("suffixes") != list(SUFFIXES) or snapshot.get("special") != {key: list(value) for key, value in SPECIAL.items()}:
        raise ValueError("风险词表规则版本不一致")
    if snapshot.get("rules_sha256") != _digest({key: value for key, value in snapshot.items() if key != "rules_sha256"}):
        raise ValueError("风险词表快照哈希不一致")
    entries = snapshot.get("entries")
    if not isinstance(entries, list) or not entries:
        raise ValueError("风险词表快照缺少关键词")
    for entry in entries:
        if not isinstance(entry, dict) or not isinstance(entry.get("keyword"), str) or not isinstance(entry.get("level"), str) or entry.get("terms") != list(dict.fromkeys((*_core_terms(entry["keyword"]), entry["keyword"]))):
            raise ValueError("风险词表快照条目非法")


def risk_matches(query: str, snapshot: Mapping[str, Any]) -> bool:
    return any(term in query for entry in snapshot["entries"] for term in entry["terms"])
